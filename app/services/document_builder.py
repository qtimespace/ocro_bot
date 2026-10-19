"""Генерация DOCX в режимах facsimile и editable (FR-07).

facsimile: каждая страница вставляется изображением (визуальный вид сохраняется, но вёрстка
не превращается в редактируемые объекты Word), затем идёт раздел с OCR-текстом по страницам.
editable: только редактируемый текст, страницы разделены разрывами, строки — отдельные абзацы.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Emu, Inches, Pt
from PIL import Image

from app.models import DocxMode, PageResult, Table

# Символы, недопустимые в XML 1.0 (python-docx на них падает).
_INVALID_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]")
# Выше этого разрешения изображение в DOCX не добавляет качества, только размер файла.
MAX_EMBED_DPI = 300
A4_WIDTH, A4_HEIGHT = Inches(8.27), Inches(11.69)
MARGIN = Inches(0.4)

VALIDATION_LABELS = {
    "not_requested": "ИИ-проверка не выполнялась (только OCR)",
    "text_checked": "текст проверен ИИ без сверки с оригиналом",
    "image_checked": "текст сверен ИИ с изображением страницы",
    "failed": "ИИ-проверка не выполнена из-за сбоя — текст НЕ проверен",
    "skipped": "ИИ-проверка пропущена — текст НЕ проверен",
}
DISCLAIMER = (
    "Текст получен автоматическим распознаванием и может содержать ошибки. "
    "ФИО, даты, суммы, номера и реквизиты сверяйте с оригиналом вручную."
)


class DocumentBuildError(RuntimeError):
    pass


def clean_text(text: str) -> str:
    return _INVALID_XML.sub("", text or "")


def _setup_section(section, landscape: bool) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = (A4_HEIGHT, A4_WIDTH) if landscape else (A4_WIDTH, A4_HEIGHT)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, MARGIN)


def _fit(width_px: int, height_px: int, max_w: int, max_h: int) -> tuple[int, int]:
    scale = min(max_w / width_px, max_h / height_px)
    return int(width_px * scale), int(height_px * scale)


def _embed_source(page: PageResult, width_emu: int, work_dir: Path) -> Path:
    """Уменьшает чрезмерно большое изображение до MAX_EMBED_DPI при заданной ширине на листе."""
    max_px = int(Emu(width_emu).inches * MAX_EMBED_DPI)
    if page.page.width <= max_px:
        return page.page.path
    target = work_dir / f"embed_{page.page.index:04d}.png"
    with Image.open(page.page.path) as im:
        ratio = max_px / im.width
        im.resize((max_px, max(1, int(im.height * ratio))), Image.Resampling.LANCZOS).save(target, format="PNG")
    return target


def _add_table(doc, table: Table) -> None:
    """Таблица Word с границами; объединённые ячейки восстанавливаются через merge."""
    word_table = doc.add_table(rows=table.rows, cols=table.cols)
    word_table.style = "Table Grid"
    for cell in table.cells:
        target = word_table.cell(cell.row, cell.col)
        if cell.rowspan > 1 or cell.colspan > 1:
            target = target.merge(word_table.cell(cell.row + cell.rowspan - 1, cell.col + cell.colspan - 1))
        target.text = clean_text(cell.text)
    doc.add_paragraph()  # отступ после таблицы, чтобы следующий текст не прилипал


def _add_page_content(doc, page: PageResult) -> None:
    if not page.blocks:
        _add_lines(doc, page.text)
        return
    for block in page.blocks:
        if isinstance(block, Table):
            _add_table(doc, block)
        else:
            p = doc.add_paragraph(clean_text(block))
            p.paragraph_format.space_after = Pt(0)


def _add_lines(doc, text: str) -> None:
    lines = clean_text(text).splitlines()
    if not any(line.strip() for line in lines):
        p = doc.add_paragraph()
        p.add_run("[текст не распознан]").italic = True
        return
    for line in lines:
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(0)


def _add_status(doc, page: PageResult) -> None:
    p = doc.add_paragraph()
    run = p.add_run(f"Статус: {VALIDATION_LABELS.get(page.validation, page.validation)}")
    run.italic = True
    run.font.size = Pt(9)


def build_docx(pages: list[PageResult], output_path: Path, mode: DocxMode = "facsimile") -> Path:
    if not pages:
        raise DocumentBuildError("No pages to build")
    if mode not in ("facsimile", "editable"):
        raise DocumentBuildError("mode must be 'facsimile' or 'editable'")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc = Document()
        doc.core_properties.title = "Распознанный документ"
        _setup_section(doc.sections[0], landscape=False)
        if mode == "facsimile":
            _build_facsimile(doc, pages, output_path.parent)
        else:
            _build_editable(doc, pages)
        doc.save(output_path)
    except DocumentBuildError:
        raise
    except Exception as exc:
        raise DocumentBuildError(f"DOCX build failed: {type(exc).__name__}") from exc
    return output_path


def _build_facsimile(doc, pages: list[PageResult], work_dir: Path) -> None:
    for i, page in enumerate(pages):
        landscape = page.page.width > page.page.height
        section = doc.sections[0] if i == 0 else doc.add_section(WD_SECTION.NEW_PAGE)
        _setup_section(section, landscape)
        max_w = section.page_width - section.left_margin - section.right_margin
        # Небольшой запас по высоте, чтобы абзац с картинкой не переносился на следующую страницу.
        max_h = section.page_height - section.top_margin - section.bottom_margin - Pt(24)
        width, height = _fit(page.page.width, page.page.height, max_w, max_h)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(_embed_source(page, width, work_dir)), width=width, height=height)

    _setup_section(doc.add_section(WD_SECTION.NEW_PAGE), landscape=False)
    doc.add_heading("Распознанный текст", level=1)
    note = doc.add_paragraph(DISCLAIMER)
    note.runs[0].italic = True
    for page in pages:
        doc.add_heading(f"Страница {page.page.index}", level=2)
        _add_status(doc, page)
        _add_page_content(doc, page)


def _build_editable(doc, pages: list[PageResult]) -> None:
    for i, page in enumerate(pages):
        if i:
            doc.add_page_break()
        doc.add_heading(f"Страница {page.page.index}", level=2)
        _add_status(doc, page)
        _add_page_content(doc, page)
    note = doc.add_paragraph(DISCLAIMER)
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(9)
