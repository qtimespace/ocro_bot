import numpy as np
from docx import Document
from PIL import Image, ImageDraw

from app.models import CellFix, OcrLine, OcrResult, PageImage, PageResult, Table, TableCell, TableFix
from app.services.ai_validator import apply_table_fixes
from app.services.document_builder import build_docx
from app.services.pipeline import Pipeline
from app.services.tables import (
    blocks_from_placeholders, blocks_to_text, blocks_with_placeholders, build_blocks, detect_tables,
    erase_table_lines, estimate_skew,
    rotate,
)
from tests.conftest import make_settings, write

XS = [100, 400, 700, 1000]
YS = [300, 380, 460, 540]


def table_page(merge_last_row: bool = True, text_bars: bool = True) -> Image.Image:
    """Страница 1200x1000: таблица 3x3 (в последней строке объединены 2 первые ячейки) и «строки текста»."""
    im = Image.new("L", (1200, 1000), 255)
    d = ImageDraw.Draw(im)
    for y in YS:
        d.line([(XS[0], y), (XS[-1], y)], fill=0, width=3)
    for i, x in enumerate(XS):
        bottom = YS[2] if merge_last_row and i == 1 else YS[-1]
        d.line([(x, YS[0]), (x, bottom)], fill=0, width=3)
    if text_bars:
        for y in (100, 140, 180, 650, 690, 730):
            d.rectangle([100, y, 900, y + 14], fill=0)
    return im


def test_detect_table_with_merged_cell():
    tables = detect_tables(np.asarray(table_page()))
    assert len(tables) == 1
    t = tables[0]
    assert (t.rows, t.cols) == (3, 3)
    merged = [c for c in t.cells if c.colspan > 1]
    assert [(c.row, c.col, c.colspan) for c in merged] == [(2, 0, 2)]
    assert len(t.cells) == 8


def test_no_table_on_plain_text_or_single_frame():
    assert detect_tables(np.asarray(table_page(text_bars=True).point(lambda v: 255))) == []
    im = Image.new("L", (1200, 1000), 255)
    ImageDraw.Draw(im).rectangle([100, 100, 1100, 900], outline=0, width=3)  # рамка страницы
    assert detect_tables(np.asarray(im)) == []


def test_skew_estimated_and_table_found_after_rotation():
    tilted = np.asarray(table_page(merge_last_row=False).rotate(-0.8, fillcolor=255))
    skew = estimate_skew(tilted)
    assert abs(skew + 0.8) < 0.25  # PIL rotate(-0.8) — поворот по часовой стрелке
    tables = detect_tables(rotate(tilted, skew))
    assert len(tables) == 1 and (tables[0].rows, tables[0].cols) == (3, 3)


def test_no_skew_on_straight_page():
    assert estimate_skew(np.asarray(table_page())) == 0.0


def frag(text, x, y):
    return OcrLine(text, box=(x, y, x + 120, y + 25))


def detected_table() -> Table:
    return detect_tables(np.asarray(table_page()))[0]


def test_build_blocks_places_fragments_in_cells_and_keeps_order():
    t = detected_table()
    ocr = OcrResult(rows=[[f] for f in [
        frag("Заголовок", 100, 100),
        frag("Наименование", 110, 310), frag("работ", 110, 340), frag("Сумма", 710, 310),
        frag("Итого", 110, 470), frag("100", 710, 470),
        frag("Подпись", 100, 650),
    ]])
    blocks = build_blocks(ocr, [t])
    assert blocks[0] == "Заголовок" and blocks[1] is t and blocks[2] == "Подпись"
    cell = {(c.row, c.col): c.text for c in t.cells}
    assert cell[(0, 0)] == "Наименование работ"
    assert cell[(0, 2)] == "Сумма"
    assert cell[(2, 0)] == "Итого" and cell[(2, 2)] == "100"
    assert "Наименование работ |  | Сумма" in blocks_to_text(blocks)


def simple_table() -> Table:
    return Table(rows=1, cols=2, cells=[TableCell(0, 0, text="Сумма"), TableCell(0, 1, text="1OO руб.")])


def test_placeholders_roundtrip_and_lost_placeholder_restored():
    t = simple_table()
    blocks = ["До", t, "После"]
    text, tables = blocks_with_placeholders(blocks)
    assert text == "До\n[[ТАБЛИЦА 1]]\nПосле" and tables == [t]
    assert blocks_from_placeholders("До!\n[[ТАБЛИЦА 1]]\nПосле", tables, blocks) == ["До!", t, "После"]
    # ИИ потерял метку — таблица возвращается на исходное место.
    assert blocks_from_placeholders("До!\nПосле", tables, blocks) == ["До!", t, "После"]


def test_apply_table_fixes():
    t = simple_table()
    assert apply_table_fixes([t], [TableFix(id=1, cells=[CellFix(r=0, c=0, text="Сумма"),
                                                         CellFix(r=0, c=1, text="100 руб.")])]) == []
    assert t.cells[1].text == "100 руб."
    # Изменение структуры отклоняется.
    notes = apply_table_fixes([t], [TableFix(id=1, cells=[CellFix(r=0, c=0, text="Сумма 100 руб.")])])
    assert "структуру" in notes[0] and t.cells[1].text == "100 руб."


def test_docx_contains_word_table_with_merge(tmp_path):
    t = Table(rows=2, cols=3, cells=[
        TableCell(0, 0, text="A"), TableCell(0, 1, text="B"), TableCell(0, 2, text="C"),
        TableCell(1, 0, colspan=2, text="Итого"), TableCell(1, 2, text="9"),
    ])
    path = write(tmp_path, "p.png", b"")
    Image.new("RGB", (600, 800), "white").save(path)
    page = PageResult(page=PageImage(1, path, 600, 800), text="", blocks=["До таблицы", t, "После"])
    doc = Document(build_docx([page], tmp_path / "r.docx", mode="editable"))
    assert len(doc.tables) == 1
    word = doc.tables[0]
    assert [c.text for c in word.rows[0].cells] == ["A", "B", "C"]
    assert word.cell(1, 0)._tc is word.cell(1, 1)._tc  # ячейки объединены
    assert word.cell(1, 0).text == "Итого" and word.cell(1, 2).text == "9"
    body = [el.tag.split("}")[1] for el in doc.element.body]
    assert body.index("tbl") > 0


class TableOcr:
    def recognize(self, image):
        return OcrResult(rows=[[f] for f in [frag("Акт", 100, 100), frag("Итого", 110, 470), frag("100", 710, 470)]])


async def test_pipeline_builds_table_blocks(tmp_path):
    settings = make_settings(tmp_path, DOCX_MODE="editable")
    work = tmp_path / "work"
    work.mkdir()
    table_page().convert("RGB").save(work / "in", format="PNG")
    result = await Pipeline(settings, TableOcr(), None).run(work / "in", work, "off")
    page = result.pages[0]
    assert page.blocks[0] == "Акт" and isinstance(page.blocks[1], Table)
    assert "Итого |  | 100" not in page.text and "Итого | 100" in page.text
    assert len(Document(result.docx_path).tables) == 1


async def test_pipeline_tables_can_be_disabled(tmp_path):
    settings = make_settings(tmp_path, DETECT_TABLES="false")
    work = tmp_path / "work"
    work.mkdir()
    table_page().convert("RGB").save(work / "in", format="PNG")
    result = await Pipeline(settings, TableOcr(), None).run(work / "in", work, "off")
    assert result.pages[0].blocks == []


def color_table_page(underline_rgb=(220, 30, 30)) -> np.ndarray:
    """Таблица 3x3 в цвете + цветное подчёркивание исправления во всю ширину средней ячейки."""
    im = table_page(merge_last_row=False).convert("RGB")
    d = ImageDraw.Draw(im)
    d.line([(XS[1] + 2, 350), (XS[2] - 2, 350)], fill=underline_rgb, width=3)
    return cv2_bgr(im)


def cv2_bgr(im: Image.Image) -> np.ndarray:
    return np.asarray(im)[:, :, ::-1].copy()


def test_colored_underline_is_not_a_row_border():
    color = color_table_page()
    gray = np.asarray(Image.fromarray(color[:, :, ::-1]).convert("L"))
    tables = detect_tables(gray, color)
    assert len(tables) == 1 and tables[0].rows == 3
    # Без цвета (серый скан) подчёркивание не отличить — это ожидаемое ограничение.
    assert detect_tables(gray)[0].rows == 4


def test_slanted_table_detected_without_deskew():
    tilted = table_page(merge_last_row=False, text_bars=False).rotate(1.5, fillcolor=255, resample=Image.BICUBIC)
    tables = detect_tables(np.asarray(tilted))
    assert len(tables) == 1 and (tables[0].rows, tables[0].cols) == (3, 3)


def test_erase_table_lines_keeps_text():
    page = table_page(merge_last_row=False)
    ImageDraw.Draw(page).rectangle([150, 320, 300, 340], fill=0)  # «текст» в ячейке
    gray = np.asarray(page)
    erased = erase_table_lines(cv2_bgr(page.convert("RGB")), detect_tables(gray))
    assert erased[YS[1], 250].min() > 200  # линия стёрта
    assert erased[330, 200].max() < 60  # текст на месте


def test_empty_edge_column_pruned():
    t = detected_table()
    # Только первые две колонки с текстом — пустая третья колонка убирается.
    ocr = OcrResult(rows=[[frag("A", 110, 310)], [frag("B", 410, 310)], [frag("C", 410, 470)]])
    blocks = build_blocks(ocr, [t])
    table = next(b for b in blocks if isinstance(b, Table))
    assert table.cols == 2
    assert {(c.row, c.col): c.text for c in table.cells}[(0, 1)] == "B"
