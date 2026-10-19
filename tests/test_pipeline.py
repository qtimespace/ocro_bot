import asyncio
from pathlib import Path

import pytest
from docx import Document

from app.services.ai_validator import AIValidationError, AIValidator
from app.services.file_parser import InputFileError
from app.services.ocr_engine import OcrError
from app.services.pipeline import Pipeline, cleanup_stale_jobs, job_workspace
from tests.conftest import FakeOcr, FakeOpenAI, image_bytes, make_settings, pdf_bytes, write


async def run(pipeline: Pipeline, tmp_path: Path, data: bytes, mode="off"):
    work = tmp_path / "work"
    work.mkdir(exist_ok=True)
    return await pipeline.run(write(work, "input", data), work, mode)


async def test_off_mode_does_not_call_ai(tmp_path, settings):
    client = FakeOpenAI([AssertionError("AI must not be called")])
    pipeline = Pipeline(settings, FakeOcr(["Hello"]), AIValidator("", "m", client=client))
    result = await run(pipeline, tmp_path, image_bytes())
    assert client.requests == []
    assert result.pages[0].validation == "not_requested"
    assert result.pages[0].text == "Hello"
    assert Document(result.docx_path).inline_shapes


async def test_multipage_pdf_order(tmp_path, settings):
    ocr = FakeOcr(["one", "two", "three"])
    result = await run(Pipeline(settings, ocr, None), tmp_path, pdf_bytes(3))
    assert [p.text for p in result.pages] == ["one", "two", "three"]
    assert [p.page.index for p in result.pages] == [1, 2, 3]
    assert ocr.calls == 3


async def test_editable_docx_mode(tmp_path):
    settings = make_settings(tmp_path, DOCX_MODE="editable")
    result = await run(Pipeline(settings, FakeOcr(["Line"]), None), tmp_path, image_bytes())
    doc = Document(result.docx_path)
    assert not doc.inline_shapes
    assert any(p.text == "Line" for p in doc.paragraphs)


async def test_no_text_produces_result_and_warning(tmp_path, settings):
    result = await run(Pipeline(settings, FakeOcr([""]), None), tmp_path, image_bytes())
    assert result.docx_path.exists()
    assert any("текст не найден" in w for w in result.warnings)


async def test_low_confidence_warning(tmp_path, settings):
    result = await run(Pipeline(settings, FakeOcr(["12345"], confidence=0.3), None), tmp_path, image_bytes())
    assert any("12345" in w and "низкая уверенность" in w for w in result.warnings)


@pytest.mark.parametrize("mode,status", [("text", "text_checked"), ("image", "image_checked")])
async def test_ai_modes(tmp_path, settings, mode, status):
    client = FakeOpenAI([{"corrected_text": "Привет", "warnings": [{"fragment": "П", "reason": "проверить"}],
                          "needs_review": True}])
    pipeline = Pipeline(settings, FakeOcr(["Првиет"]), AIValidator("", "m", client=client))
    result = await run(pipeline, tmp_path, image_bytes(), mode)
    assert result.pages[0].text == "Привет"
    assert result.pages[0].validation == status
    assert any("«П» — проверить" in w for w in result.warnings)
    assert any("ручную проверку" in w for w in result.warnings)
    has_image = "image_url" in str(client.requests[0]["messages"])
    assert has_image == (mode == "image")


async def test_ai_failure_returns_unverified_result(tmp_path):
    settings = make_settings(tmp_path, AI_CONCURRENCY="1")
    client = FakeOpenAI([ConnectionError("down")])
    pipeline = Pipeline(settings, FakeOcr(["a", "b"]), AIValidator("", "m", client=client))
    result = await run(pipeline, tmp_path, pdf_bytes(2), "image")
    assert [p.validation for p in result.pages] == ["failed", "skipped"]
    assert [p.text for p in result.pages] == ["a", "b"]
    assert len(client.requests) == 1  # после сбоя остальные страницы не отправляются
    assert result.ai_failed
    assert any("не выполнена" in w for w in result.warnings)
    assert "НЕ проверен" in "\n".join(p.text for p in Document(result.docx_path).paragraphs)


async def test_ai_mode_without_key_rejected(tmp_path, settings):
    with pytest.raises(AIValidationError):
        await run(Pipeline(settings, FakeOcr(), None), tmp_path, image_bytes(), "text")


async def test_ocr_error_propagates(tmp_path, settings):
    with pytest.raises(OcrError):
        await run(Pipeline(settings, FakeOcr(error=OcrError("x")), None), tmp_path, image_bytes())


async def test_invalid_file_rejected_before_ocr(tmp_path, settings):
    ocr = FakeOcr()
    with pytest.raises(InputFileError):
        await run(Pipeline(settings, ocr, None), tmp_path, b"not an image")
    assert ocr.calls == 0


def test_workspace_removed_on_success_and_error(tmp_path):
    root = tmp_path / "temp"
    with job_workspace(root) as work:
        (work / "file.bin").write_bytes(b"data")
    assert not work.exists()
    with pytest.raises(RuntimeError):
        with job_workspace(root) as work2:
            (work2 / "file.bin").write_bytes(b"data")
            raise RuntimeError("fail")
    assert not work2.exists()
    assert list(root.iterdir()) == []


def test_cleanup_stale_jobs(tmp_path):
    import os
    import time

    root = tmp_path / "temp"
    old, fresh, other = root / "job_old", root / "job_fresh", root / "keep"
    for d in (old, fresh, other):
        d.mkdir(parents=True)
    past = time.time() - 7200
    os.utime(old, (past, past))
    os.utime(other, (past, past))
    assert cleanup_stale_jobs(root) == 1
    assert not old.exists() and fresh.exists() and other.exists()


async def test_multiple_inputs_combined_and_reordered(tmp_path, settings):
    work = tmp_path / "work"
    work.mkdir()
    inputs = [write(work, f"in{i}", image_bytes()) for i in range(3)]
    ocr = FakeOcr(["B\n^Стр. 2", "C\n^Стр. 3", "A\n^Стр. 1"])
    result = await Pipeline(settings, ocr, None).run(inputs, work, "off")
    assert [p.text.split("\n")[0] for p in result.pages] == ["A", "B", "C"]
    assert [p.page.index for p in result.pages] == [1, 2, 3]
    assert any("восстановлен" in w for w in result.warnings)


async def test_multipage_file_not_split_when_reordering(tmp_path):
    settings = make_settings(tmp_path, PDF_DPI="72")  # высота страницы PDF ~842 px — колонтитул фейка внизу
    work = tmp_path / "work"
    work.mkdir()
    inputs = [write(work, "a", image_bytes()), write(work, "b", pdf_bytes(2))]
    # Файл 1 (стр. 3) и PDF из двух страниц (стр. 1, 2): PDF переносится целиком вперёд.
    ocr = FakeOcr(["Z\n^3", "X\n^1", "Y\n^2"])
    result = await Pipeline(settings, ocr, None).run(inputs, work, "off")
    assert [p.text.split("\n")[0] for p in result.pages] == ["X", "Y", "Z"]


async def test_batch_page_limit(tmp_path):
    settings = make_settings(tmp_path, MAX_PDF_PAGES="2")
    work = tmp_path / "work"
    work.mkdir()
    inputs = [write(work, "a", pdf_bytes(2)), write(work, "b", image_bytes())]
    with pytest.raises(InputFileError, match="больше 2"):
        await Pipeline(settings, FakeOcr(), None).run(inputs, work, "off")


async def test_ai_pages_checked_in_parallel(tmp_path):
    settings = make_settings(tmp_path, AI_CONCURRENCY="4")
    active = {"now": 0, "max": 0}

    class SlowClient(FakeOpenAI):
        async def _create(self, **kwargs):
            active["now"] += 1
            active["max"] = max(active["max"], active["now"])
            await asyncio.sleep(0.05)
            active["now"] -= 1
            return await super()._create(**kwargs)

    client = SlowClient([{"corrected_text": "ok"}])
    pipeline = Pipeline(settings, FakeOcr(["ok"]), AIValidator("", "m", client=client))
    work = tmp_path / "work"
    work.mkdir()
    inputs = [write(work, f"in{i}", image_bytes()) for i in range(4)]
    result = await pipeline.run(inputs, work, "text")
    assert active["max"] > 1
    assert all(p.validation == "text_checked" for p in result.pages)
