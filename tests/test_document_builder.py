from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT

from app.models import PageImage, PageResult
from app.services.document_builder import DocumentBuildError, build_docx
from tests.conftest import image_bytes, write


def page(tmp_path: Path, index: int, text: str, size=(600, 800), validation="not_requested") -> PageResult:
    path = write(tmp_path, f"p{index}.png", image_bytes("PNG", size=size))
    return PageResult(page=PageImage(index=index, path=path, width=size[0], height=size[1]), text=text,
                      validation=validation)


def all_text(doc) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def test_facsimile_has_images_in_order_and_text(tmp_path):
    pages = [page(tmp_path, 1, "Hello\nПривет"), page(tmp_path, 2, "Second page")]
    out = build_docx(pages, tmp_path / "result.docx", mode="facsimile")
    doc = Document(out)
    assert len(doc.inline_shapes) == 2
    text = all_text(doc)
    assert "Распознанный текст" in text
    assert text.index("Страница 1") < text.index("Hello") < text.index("Страница 2") < text.index("Second page")
    assert any(p.text == "Привет" for p in doc.paragraphs)


def test_facsimile_keeps_aspect_ratio_and_fits_page(tmp_path):
    out = build_docx([page(tmp_path, 1, "x", size=(1000, 2000))], tmp_path / "r.docx")
    doc = Document(out)
    shape = doc.inline_shapes[0]
    assert abs(shape.width / shape.height - 0.5) < 0.01
    section = doc.sections[0]
    assert shape.height <= section.page_height - section.top_margin - section.bottom_margin


def test_facsimile_landscape_page(tmp_path):
    out = build_docx([page(tmp_path, 1, "wide", size=(1200, 600))], tmp_path / "r.docx")
    assert Document(out).sections[0].orientation == WD_ORIENT.LANDSCAPE


def test_large_image_downscaled_for_embedding(tmp_path):
    big = page(tmp_path, 1, "x", size=(6000, 8000))
    out = build_docx([big], tmp_path / "r.docx")
    assert out.stat().st_size < big.page.path.stat().st_size * 2
    assert len(Document(out).inline_shapes) == 1


def test_editable_lines_and_page_breaks(tmp_path):
    out = build_docx([page(tmp_path, 1, "Line one\nLine two"), page(tmp_path, 2, "Page two")],
                     tmp_path / "r.docx", mode="editable")
    doc = Document(out)
    assert len(doc.inline_shapes) == 0
    texts = [p.text for p in doc.paragraphs]
    assert "Line one" in texts and "Line two" in texts and "Page two" in texts
    assert 'w:br w:type="page"' in doc.element.xml


def test_empty_text_marked(tmp_path):
    doc = Document(build_docx([page(tmp_path, 1, "")], tmp_path / "r.docx", mode="editable"))
    assert "[текст не распознан]" in all_text(doc)


def test_invalid_xml_chars_removed(tmp_path):
    doc = Document(build_docx([page(tmp_path, 1, "A\x00B\x0cC\x1f")], tmp_path / "r.docx", mode="editable"))
    assert "ABC" in all_text(doc)


def test_validation_status_written(tmp_path):
    doc = Document(build_docx([page(tmp_path, 1, "x", validation="failed")], tmp_path / "r.docx"))
    assert "НЕ проверен" in all_text(doc)


def test_invalid_mode_and_empty_pages(tmp_path):
    with pytest.raises(DocumentBuildError):
        build_docx([], tmp_path / "r.docx")
    with pytest.raises(DocumentBuildError):
        build_docx([page(tmp_path, 1, "x")], tmp_path / "r.docx", mode="pdf")  # type: ignore[arg-type]
